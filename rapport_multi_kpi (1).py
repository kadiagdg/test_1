"""
rapport_multi_kpi.py
=====================
Assemble PLUSIEURS DataFrames de KPI déjà calculés dans un unique classeur
Excel journalier, avec un sommaire de navigation et une feuille par KPI.

Ce script ne calcule AUCUN KPI : il prend en entrée les DataFrames que tu as
déjà produits ailleurs (tes fonctions de calcul, requêtes SQL, exports Kobo,
etc.) et se charge uniquement de la mise en forme et de l'assemblage.

Dépendances : pandas, openpyxl
    pip install pandas openpyxl --break-system-packages
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ----------------------------------------------------------------------
# CONFIGURATION / PALETTE
# ----------------------------------------------------------------------

OUTPUT_DIR = Path("rapports_kpi")
OUTPUT_DIR.mkdir(exist_ok=True)

COLOR_HEADER_BG = "1F4E78"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_OK = "C6EFCE"
COLOR_WARN = "FFEB9C"
COLOR_KO = "FFC7CE"

MAX_SHEET_NAME_LEN = 31  # limite Excel
CARACTERES_INTERDITS = r'[:\\/?*\[\]]'


@dataclass
class SeuilKPI:
    """Seuil d'alerte appliqué à une colonne numérique d'un KPI."""
    colonne: str
    seuil_ko: float
    seuil_warn: float
    sens: str = "min"   # "min" = plus haut = mieux (ex: taux) / "max" = plus haut = pire (ex: NPL)


@dataclass
class DefinitionKPI:
    """Emballe un DataFrame déjà calculé avec ses métadonnées de présentation."""
    nom: str
    df: pd.DataFrame
    seuils: list[SeuilKPI] = field(default_factory=list)
    colonne_cle: str | None = None  # colonne résumée (moyenne) sur la page Sommaire


# ----------------------------------------------------------------------
# UTILITAIRES DE MISE EN FORME (inchangés — pas besoin d'y toucher)
# ----------------------------------------------------------------------

def nom_feuille_valide(nom: str) -> str:
    nom_propre = re.sub(CARACTERES_INTERDITS, "-", nom)
    return nom_propre[:MAX_SHEET_NAME_LEN]


def appliquer_bordures_et_largeurs(ws: Worksheet, df: pd.DataFrame, ligne_depart: int) -> None:
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=ligne_depart + 1, max_row=ligne_depart + df.shape[0], max_col=df.shape[1]):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    for col_idx, col_name in enumerate(df.columns, start=1):
        longueur_max = max(df[col_name].astype(str).map(len).max(), len(str(col_name)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(longueur_max + 4, 40)


def activer_filtre_auto(ws: Worksheet, df: pd.DataFrame, ligne_depart: int) -> None:
    derniere_col = get_column_letter(df.shape[1])
    ws.auto_filter.ref = f"A{ligne_depart}:{derniere_col}{ligne_depart + df.shape[0]}"


def appliquer_seuils(ws: Worksheet, df: pd.DataFrame, seuils: list[SeuilKPI], ligne_depart: int) -> None:
    fill_ok = PatternFill(start_color=COLOR_OK, end_color=COLOR_OK, fill_type="solid")
    fill_warn = PatternFill(start_color=COLOR_WARN, end_color=COLOR_WARN, fill_type="solid")
    fill_ko = PatternFill(start_color=COLOR_KO, end_color=COLOR_KO, fill_type="solid")

    for seuil in seuils:
        if seuil.colonne not in df.columns:
            continue
        col_letter = get_column_letter(df.columns.get_loc(seuil.colonne) + 1)
        plage = f"{col_letter}{ligne_depart + 1}:{col_letter}{ligne_depart + df.shape[0]}"

        if seuil.sens == "min":
            ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_ko)], fill=fill_ko))
            ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_ko), str(seuil.seuil_warn)], fill=fill_warn))
            ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_warn)], fill=fill_ok))
        else:
            ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_ko)], fill=fill_ko))
            ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_warn), str(seuil.seuil_ko)], fill=fill_warn))
            ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_warn)], fill=fill_ok))


def construire_feuille_kpi(wb: Workbook, kpi: DefinitionKPI) -> str:
    """Crée une feuille Excel formatée à partir d'un DataFrame déjà calculé."""
    nom_feuille = nom_feuille_valide(kpi.nom)
    nom_final, i = nom_feuille, 1
    while nom_final in wb.sheetnames:
        suffixe = f"_{i}"
        nom_final = nom_feuille[: MAX_SHEET_NAME_LEN - len(suffixe)] + suffixe
        i += 1

    ws = wb.create_sheet(nom_final)

    ws.cell(row=1, column=1, value="⇦ Retour au sommaire").hyperlink = "#Sommaire!A1"
    ws.cell(row=1, column=1).font = Font(color="0563C1", underline="single", size=9)
    ws.cell(row=2, column=1, value=f"{kpi.nom} — {len(kpi.df)} ligne(s)").font = Font(bold=True, size=13)

    ligne_tableau = 4
    for j, col_name in enumerate(kpi.df.columns, start=1):
        ws.cell(row=ligne_tableau, column=j, value=str(col_name))
    for i, row in enumerate(kpi.df.itertuples(index=False), start=ligne_tableau + 1):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)

    for col_idx in range(1, kpi.df.shape[1] + 1):
        cell = ws.cell(row=ligne_tableau, column=col_idx)
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.font = Font(color=COLOR_HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = f"A{ligne_tableau + 1}"
    appliquer_bordures_et_largeurs(ws, kpi.df, ligne_depart=ligne_tableau)
    activer_filtre_auto(ws, kpi.df, ligne_depart=ligne_tableau)
    if kpi.seuils:
        appliquer_seuils(ws, kpi.df, kpi.seuils, ligne_depart=ligne_tableau)

    return nom_final


def construire_feuille_sommaire(wb: Workbook, kpis: list[DefinitionKPI], noms_feuilles: dict[str, str], date_rapport: dt.date) -> None:
    ws = wb.create_sheet("Sommaire", 0)

    titre = ws.cell(row=1, column=1, value=f"Rapport KPI journalier — {date_rapport.strftime('%d/%m/%Y')}")
    titre.font = Font(size=16, bold=True, color=COLOR_HEADER_BG)
    ws.merge_cells("A1:D1")

    entetes = ["KPI", "Nb lignes", "Valeur clé (moy.)", "Accès"]
    ligne_entete = 3
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=ligne_entete, column=j, value=e)
        c.font = Font(bold=True, color=COLOR_HEADER_FONT)
        c.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")

    for i, kpi in enumerate(kpis, start=ligne_entete + 1):
        nom_feuille = noms_feuilles[kpi.nom]
        ws.cell(row=i, column=1, value=kpi.nom)
        ws.cell(row=i, column=2, value=len(kpi.df))
        if kpi.colonne_cle and kpi.colonne_cle in kpi.df.columns:
            ws.cell(row=i, column=3, value=round(kpi.df[kpi.colonne_cle].mean(), 2))
        lien = ws.cell(row=i, column=4, value="Voir le détail →")
        lien.hyperlink = f"#'{nom_feuille}'!A1"
        lien.font = Font(color="0563C1", underline="single")

    for col, largeur in zip("ABCD", [38, 12, 20, 18]):
        ws.column_dimensions[col].width = largeur


# ----------------------------------------------------------------------
# POINT D'ENTRÉE PRINCIPAL — c'est la fonction que tu appelles depuis ton code
# ----------------------------------------------------------------------

def generer_rapport_multi_kpi(kpis: list[DefinitionKPI], date_rapport: dt.date | None = None) -> Path:
    """
    Assemble une liste de KPI (déjà calculés) dans un classeur Excel unique.

    Paramètres
    ----------
    kpis : liste de DefinitionKPI, un par tableau de sortie
    date_rapport : date affichée dans le rapport (par défaut : aujourd'hui)
    """
    date_rapport = date_rapport or dt.date.today()
    wb = Workbook()
    wb.remove(wb.active)

    noms_feuilles: dict[str, str] = {}
    for kpi in kpis:
        noms_feuilles[kpi.nom] = construire_feuille_kpi(wb, kpi)

    construire_feuille_sommaire(wb, kpis, noms_feuilles, date_rapport)

    fichier_sortie = OUTPUT_DIR / f"rapport_multi_kpi_{date_rapport.isoformat()}.xlsx"
    wb.save(fichier_sortie)
    return fichier_sortie


# ----------------------------------------------------------------------
# ============  ICI : TU BRANCHES TES VRAIS DATAFRAMES  ================
# ----------------------------------------------------------------------
#
# Remplace ce bloc par tes df déjà calculés (issus de tes fonctions,
# de tes requêtes SQL, de ton pipeline, etc.). Chaque DefinitionKPI
# correspond à un tableau de sortie, quelle que soit sa taille (3 à
# 100+ lignes) — aucune adaptation nécessaire selon la taille.

def main() -> None:
    # Exemple : tu as déjà ces variables en sortie de tes calculs
    # df_taux_transformation = calculer_taux_transformation(...)
    # df_taux_npl            = calculer_taux_npl(...)
    # df_alertes_caisse      = detecter_depassements_caisse(...)
    # ... etc. jusqu'à tes 15 DataFrames

    kpis = [
        DefinitionKPI(
            nom="Taux de transformation",
            df=df_taux_transformation,
            seuils=[SeuilKPI(colonne="Taux (%)", seuil_ko=65, seuil_warn=80, sens="min")],
            colonne_cle="Taux (%)",
        ),
        DefinitionKPI(
            nom="Taux NPL",
            df=df_taux_npl,
            seuils=[SeuilKPI(colonne="Taux NPL (%)", seuil_ko=6.0, seuil_warn=4.0, sens="max")],
            colonne_cle="Taux NPL (%)",
        ),
        DefinitionKPI(
            nom="Alertes dépassement plafond caisse",
            df=df_alertes_caisse,
            colonne_cle="Dépassement (M FCFA)",
        ),
        # DefinitionKPI(nom="...", df=df_..., seuils=[...], colonne_cle="..."),
        # -> répète pour chacun de tes 15 KPI
    ]

    fichier = generer_rapport_multi_kpi(kpis)
    print(f"Rapport généré : {fichier.resolve()} ({len(kpis)} KPI)")


if __name__ == "__main__":
    main()

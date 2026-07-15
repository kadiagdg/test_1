"""
rapport_multi_kpi.py
=====================
Génère un rapport Excel journalier regroupant PLUSIEURS KPI (ex: 15),
chacun étant un tableau de taille variable (3 à 100+ lignes).

Structure du classeur produit :
    - Feuille "Sommaire"  : vue d'ensemble + liens hypertexte vers chaque KPI
    - 1 feuille par KPI   : tableau structuré (filtre, en-tête figé, mise en
                            forme conditionnelle optionnelle) + lien de retour

Dépendances : pandas, openpyxl
    pip install pandas openpyxl --break-system-packages
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ----------------------------------------------------------------------
# CONFIGURATION / PALETTE
# ----------------------------------------------------------------------

OUTPUT_DIR = Path("rapports_kpi")
OUTPUT_DIR.mkdir(exist_ok=True)

COLOR_HEADER_BG = "1F4E78"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_OK = "C6EFCE"
COLOR_OK_FONT = "006100"
COLOR_WARN = "FFEB9C"
COLOR_WARN_FONT = "9C6500"
COLOR_KO = "FFC7CE"
COLOR_KO_FONT = "9C0006"

MAX_SHEET_NAME_LEN = 31  # limite Excel
CARACTERES_INTERDITS = r'[:\\/?*\[\]]'


@dataclass
class SeuilKPI:
    colonne: str
    seuil_ko: float
    seuil_warn: float
    sens: str = "min"


@dataclass
class DefinitionKPI:
    """Décrit un KPI à publier : son nom, son tableau, et ses seuils d'alerte."""
    nom: str
    df: pd.DataFrame
    seuils: list[SeuilKPI] = field(default_factory=list)
    colonne_cle: str | None = None  # colonne utilisée pour le résumé en page Sommaire


# ----------------------------------------------------------------------
# DONNÉES D'EXEMPLE — remplace ceci par tes 15 vrais calculs de KPI
# ----------------------------------------------------------------------

def generer_kpis_exemple() -> list[DefinitionKPI]:
    import random

    random.seed(7)

    def table_agences(n, col_val, low, high):
        return pd.DataFrame({
            "Entité": [f"Agence {i:03d}" for i in range(1, n + 1)],
            col_val: [round(random.uniform(low, high), 2) for _ in range(n)],
            "Date": [dt.date.today()] * n,
        })

    kpis = [
        DefinitionKPI(
            "Taux de transformation",
            table_agences(45, "Taux (%)", 55, 98),
            seuils=[SeuilKPI("Taux (%)", 65, 80, "min")],
            colonne_cle="Taux (%)",
        ),
        DefinitionKPI(
            "Taux NPL",
            table_agences(45, "Taux NPL (%)", 1.5, 9.0),
            seuils=[SeuilKPI("Taux NPL (%)", 6.0, 4.0, "max")],
            colonne_cle="Taux NPL (%)",
        ),
        DefinitionKPI(
            "Alertes dépassement plafond caisse",
            table_agences(7, "Dépassement (M FCFA)", 0.5, 12),
            colonne_cle="Dépassement (M FCFA)",
        ),
        DefinitionKPI(
            "NPS clients",
            table_agences(45, "NPS", -20, 80),
            seuils=[SeuilKPI("NPS", 0, 40, "min")],
            colonne_cle="NPS",
        ),
        DefinitionKPI(
            "Dossiers en attente > 48h",
            table_agences(3, "Nb dossiers", 1, 15),
            colonne_cle="Nb dossiers",
        ),
        # ... complète jusqu'à 15 en suivant le même schéma
    ]
    return kpis


# ----------------------------------------------------------------------
# UTILITAIRES
# ----------------------------------------------------------------------

def nom_feuille_valide(nom: str) -> str:
    """Nettoie un nom de KPI pour qu'il soit un nom de feuille Excel valide."""
    nom_propre = re.sub(CARACTERES_INTERDITS, "-", nom)
    return nom_propre[:MAX_SHEET_NAME_LEN]


def styliser_entete(ws: Worksheet, n_colonnes: int) -> None:
    header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    header_font = Font(color=COLOR_HEADER_FONT, bold=True, size=11)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx in range(1, n_colonnes + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26


def appliquer_bordures_et_largeurs(ws: Worksheet, df: pd.DataFrame, ligne_depart: int = 1) -> None:
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=ligne_depart + 1, max_row=ligne_depart + df.shape[0], max_col=df.shape[1]):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    for col_idx, col_name in enumerate(df.columns, start=1):
        longueur_max = max(df[col_name].astype(str).map(len).max(), len(col_name))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(longueur_max + 4, 40)


def appliquer_mise_en_forme_conditionnelle(ws: Worksheet, df: pd.DataFrame, seuils: list[SeuilKPI]) -> None:
    n_lignes = df.shape[0]
    for seuil in seuils:
        if seuil.colonne not in df.columns:
            continue
        col_idx = df.columns.get_loc(seuil.colonne) + 1
        col_letter = get_column_letter(col_idx)
        plage = f"{col_letter}2:{col_letter}{n_lignes + 1}"

        fill_ok = PatternFill(start_color=COLOR_OK, end_color=COLOR_OK, fill_type="solid")
        font_ok = Font(color=COLOR_OK_FONT)
        fill_warn = PatternFill(start_color=COLOR_WARN, end_color=COLOR_WARN, fill_type="solid")
        font_warn = Font(color=COLOR_WARN_FONT)
        fill_ko = PatternFill(start_color=COLOR_KO, end_color=COLOR_KO, fill_type="solid")
        font_ko = Font(color=COLOR_KO_FONT)

        if seuil.sens == "min":
            ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_ko)], fill=fill_ko, font=font_ko))
            ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_ko), str(seuil.seuil_warn)], fill=fill_warn, font=font_warn))
            ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_warn)], fill=fill_ok, font=font_ok))
        else:
            ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_ko)], fill=fill_ko, font=font_ko))
            ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_warn), str(seuil.seuil_ko)], fill=fill_warn, font=font_warn))
            ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_warn)], fill=fill_ok, font=font_ok))


def activer_filtre_auto(ws: Worksheet, df: pd.DataFrame, ligne_depart: int = 1) -> None:
    derniere_col = get_column_letter(df.shape[1])
    ws.auto_filter.ref = f"A{ligne_depart}:{derniere_col}{ligne_depart + df.shape[0]}"


# ----------------------------------------------------------------------
# CONSTRUCTION D'UNE FEUILLE KPI INDIVIDUELLE
# ----------------------------------------------------------------------

def construire_feuille_kpi(wb: Workbook, kpi: DefinitionKPI) -> str:
    nom_feuille = nom_feuille_valide(kpi.nom)
    # gère les doublons de noms après troncature
    nom_final, i = nom_feuille, 1
    while nom_final in wb.sheetnames:
        suffixe = f"_{i}"
        nom_final = nom_feuille[: MAX_SHEET_NAME_LEN - len(suffixe)] + suffixe
        i += 1

    ws = wb.create_sheet(nom_final)

    # Titre + lien retour vers le sommaire
    ws.cell(row=1, column=1, value=f"⇦ Retour au sommaire").hyperlink = "#Sommaire!A1"
    ws.cell(row=1, column=1).font = Font(color="0563C1", underline="single", size=9)

    ws.cell(row=2, column=1, value=f"{kpi.nom} — {len(kpi.df)} ligne(s)").font = Font(bold=True, size=13)

    ligne_tableau = 4  # ligne où démarre l'en-tête du tableau
    for j, col_name in enumerate(kpi.df.columns, start=1):
        ws.cell(row=ligne_tableau, column=j, value=col_name)
    for i, row in enumerate(kpi.df.itertuples(index=False), start=ligne_tableau + 1):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)

    # Applique le style à partir de la ligne d'en-tête réelle
    for col_idx in range(1, kpi.df.shape[1] + 1):
        cell = ws.cell(row=ligne_tableau, column=col_idx)
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.font = Font(color=COLOR_HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = f"A{ligne_tableau + 1}"
    appliquer_bordures_et_largeurs(ws, kpi.df, ligne_depart=ligne_tableau)
    activer_filtre_auto(ws, kpi.df, ligne_depart=ligne_tableau)

    if kpi.seuils:
        # décale les seuils vers la bonne plage de lignes (ligne_tableau au lieu de 1)
        for seuil in kpi.seuils:
            if seuil.colonne not in kpi.df.columns:
                continue
            col_letter = get_column_letter(kpi.df.columns.get_loc(seuil.colonne) + 1)
            plage = f"{col_letter}{ligne_tableau + 1}:{col_letter}{ligne_tableau + kpi.df.shape[0]}"
            fill_ok = PatternFill(start_color=COLOR_OK, end_color=COLOR_OK, fill_type="solid")
            fill_warn = PatternFill(start_color=COLOR_WARN, end_color=COLOR_WARN, fill_type="solid")
            fill_ko = PatternFill(start_color=COLOR_KO, end_color=COLOR_KO, fill_type="solid")
            if seuil.sens == "min":
                ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_ko)], fill=fill_ko))
                ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_ko), str(seuil.seuil_warn)], fill=fill_warn))
                ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_warn)], fill=fill_ok))
            else:
                ws.conditional_formatting.add(plage, CellIsRule(operator="greaterThan", formula=[str(seuil.seuil_ko)], fill=fill_ko))
                ws.conditional_formatting.add(plage, CellIsRule(operator="between", formula=[str(seuil.seuil_warn), str(seuil.seuil_ko)], fill=fill_warn))
                ws.conditional_formatting.add(plage, CellIsRule(operator="lessThan", formula=[str(seuil.seuil_warn)], fill=fill_ok))

    return nom_final


# ----------------------------------------------------------------------
# CONSTRUCTION DE LA FEUILLE SOMMAIRE
# ----------------------------------------------------------------------

def construire_feuille_sommaire(wb: Workbook, kpis: list[DefinitionKPI], noms_feuilles: dict[str, str], date_rapport: dt.date) -> None:
    ws = wb.create_sheet("Sommaire", 0)

    titre = ws.cell(row=1, column=1, value=f"Rapport KPI journalier — {date_rapport.strftime('%d/%m/%Y')}")
    titre.font = Font(size=16, bold=True, color=COLOR_HEADER_BG)
    ws.merge_cells("A1:D1")

    entetes = ["KPI", "Nb lignes", "Valeur clé (moy.)", "Accès"]
    ligne_entete = 3
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=ligne_entete, column=j, value=e)
        c.font = Font(bold=True, color=COLOR_HEADER_FONT)
        c.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")

    for i, kpi in enumerate(kpis, start=ligne_entete + 1):
        nom_feuille = noms_feuilles[kpi.nom]
        ws.cell(row=i, column=1, value=kpi.nom)
        ws.cell(row=i, column=2, value=len(kpi.df))
        if kpi.colonne_cle and kpi.colonne_cle in kpi.df.columns:
            ws.cell(row=i, column=3, value=round(kpi.df[kpi.colonne_cle].mean(), 2))
        lien = ws.cell(row=i, column=4, value="Voir le détail →")
        lien.hyperlink = f"#'{nom_feuille}'!A1"
        lien.font = Font(color="0563C1", underline="single")

    for col, largeur in zip("ABCD", [38, 12, 20, 18]):
        ws.column_dimensions[col].width = largeur

    ws.cell(row=1, column=1).hyperlink = None  # ancre "#Sommaire!A1" utilisée par les liens retour


# ----------------------------------------------------------------------
# ORCHESTRATION
# ----------------------------------------------------------------------

def generer_rapport_multi_kpi(kpis: list[DefinitionKPI], date_rapport: dt.date | None = None) -> Path:
    date_rapport = date_rapport or dt.date.today()
    wb = Workbook()
    wb.remove(wb.active)

    noms_feuilles: dict[str, str] = {}
    for kpi in kpis:
        nom_final = construire_feuille_kpi(wb, kpi)
        noms_feuilles[kpi.nom] = nom_final

    construire_feuille_sommaire(wb, kpis, noms_feuilles, date_rapport)

    fichier_sortie = OUTPUT_DIR / f"rapport_multi_kpi_{date_rapport.isoformat()}.xlsx"
    wb.save(fichier_sortie)
    return fichier_sortie


def main() -> None:
    kpis = generer_kpis_exemple()
    fichier = generer_rapport_multi_kpi(kpis)
    print(f"Rapport généré : {fichier.resolve()} ({len(kpis)} KPI)")


if __name__ == "__main__":
    main()
